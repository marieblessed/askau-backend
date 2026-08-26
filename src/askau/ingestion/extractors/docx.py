"""DOCX extraction (FR-012).

Word is the easy case and the one to get exactly right: paragraph styles carry
real heading levels, so `heading_path` is read rather than inferred. Where a PDF
extractor guesses from font size, this one knows.

Tables are emitted as rows rather than flattened into prose. A per-diem table
rendered as "Region Rate Continental 180 Intercontinental 250" reads as
nonsense and retrieves badly; one row per block keeps the pairing intact.
"""

from __future__ import annotations

import re

from askau.domain.knowledge import ExtractedBlock, ExtractionResult

_HEADING_STYLE = re.compile(r"^Heading\s*(\d+)$", re.IGNORECASE)
_SECTION_REF = re.compile(r"^\s*(\d+(?:\.\d+){0,3})[.)]?\s+")


class DocxExtractor:
    def extract(self, data: bytes, filename: str) -> ExtractionResult:
        try:
            import io

            from docx import Document
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "python-docx is required for DOCX extraction: pip install '.[ingestion]'"
            ) from exc

        document = Document(io.BytesIO(data))
        blocks: list[ExtractedBlock] = []
        heading_path: tuple[str, ...] = ()
        offset = 0

        for paragraph in document.paragraphs:
            text = " ".join(paragraph.text.split())
            if not text:
                continue

            level = _heading_level(paragraph)
            if level is not None:
                heading_path = (*heading_path[: level - 1], text)
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        page=None,  # Word has no fixed pagination until rendered
                        heading_path=heading_path,
                        section_ref=_section_ref(text),
                        is_heading=True,
                        char_start=offset,
                        char_end=offset + len(text),
                    )
                )
            else:
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        page=None,
                        heading_path=heading_path,
                        section_ref=_section_ref(text) if heading_path else None,
                        char_start=offset,
                        char_end=offset + len(text),
                    )
                )
            offset += len(text) + 1

        for table_index, table in enumerate(document.tables, start=1):
            header: list[str] = []
            for row_index, row in enumerate(table.rows):
                cells = [" ".join(c.text.split()) for c in row.cells]
                if not any(cells):
                    continue
                if row_index == 0:
                    header = cells
                    continue
                # Pair each value with its column heading, so a retrieved row
                # still says what its numbers mean.
                pairs = [f"{h}: {v}" for h, v in zip(header, cells, strict=False) if v] or cells
                text = " · ".join(pairs)
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        page=None,
                        heading_path=(*heading_path, f"Table {table_index}"),
                        char_start=offset,
                        char_end=offset + len(text),
                    )
                )
                offset += len(text) + 1

        return ExtractionResult(blocks=tuple(blocks), page_count=None)


def _heading_level(paragraph: object) -> int | None:
    style = getattr(getattr(paragraph, "style", None), "name", "") or ""
    match = _HEADING_STYLE.match(style)
    if match:
        return min(int(match.group(1)), 4)
    if style.lower() in {"title", "subtitle"}:
        return 1
    return None


def _section_ref(text: str) -> str | None:
    match = _SECTION_REF.match(text)
    return match.group(1) if match else None

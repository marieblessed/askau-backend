"""PPTX extraction (FR-012).

Slides map cleanly onto the model already in place: the slide number is the
page, and the title placeholder is the heading. Speaker notes are included —
in AUC decks the substance is often in the notes while the slide carries only
bullet fragments.
"""

from __future__ import annotations

from askau.domain.knowledge import ExtractedBlock, ExtractionResult


class PptxExtractor:
    def extract(self, data: bytes, filename: str) -> ExtractionResult:
        try:
            import io

            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "python-pptx is required for PPTX extraction: pip install '.[ingestion]'"
            ) from exc

        presentation = Presentation(io.BytesIO(data))
        blocks: list[ExtractedBlock] = []
        offset = 0

        for number, slide in enumerate(presentation.slides, start=1):
            title = _title_of(slide) or f"Slide {number}"
            body: list[str] = []

            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = " ".join(shape.text_frame.text.split())
                if text and text != title:
                    body.append(text)

            if slide.has_notes_slide:
                notes = " ".join(slide.notes_slide.notes_text_frame.text.split())
                if notes:
                    body.append(f"Speaker notes: {notes}")

            blocks.append(
                ExtractedBlock(
                    text=title,
                    page=number,
                    heading_path=(title,),
                    is_heading=True,
                    char_start=offset,
                    char_end=offset + len(title),
                )
            )
            offset += len(title) + 1

            for text in body:
                blocks.append(
                    ExtractedBlock(
                        text=text,
                        page=number,
                        heading_path=(title,),
                        char_start=offset,
                        char_end=offset + len(text),
                    )
                )
                offset += len(text) + 1

        return ExtractionResult(blocks=tuple(blocks), page_count=len(presentation.slides))


def _title_of(slide: object) -> str | None:
    placeholder = getattr(getattr(slide, "shapes", None), "title", None)
    if placeholder is None:
        return None
    text = " ".join(placeholder.text.split())
    return text or None
